"""測定Excelの読込と電圧・電流の変換。GUIと一括計算で共用する。"""
from dataclasses import dataclass
import math
from pathlib import Path

import numpy as np
import openpyxl

from calculation import RESISTANCE_OHM, correct_trace

ROOT = Path(__file__).resolve().parent


@dataclass
class Measurement:
    sheet: str
    values: dict
    original_formulas: dict
    all_rows: list[int]
    all_voltage: np.ndarray
    all_current: np.ndarray
    total: np.ndarray
    applicable: np.ndarray
    selection: np.ndarray

    @property
    def rows(self):
        return np.asarray(self.all_rows)[self.selection]

    @property
    def voltage(self):
        return self.all_voltage[self.selection]

    @property
    def current(self):
        return self.all_current[self.selection]


def load_measurement(source: Path, config: dict) -> Measurement:
    """セルの保存済み計算結果に依存せず、F/H列の生測定値から再計算。"""
    book = openpyxl.load_workbook(source, read_only=True, data_only=False)
    try:
        if config["sheet"] not in book.sheetnames:
            raise ValueError(f"対象シートがありません: {config['sheet']}")
        sheet = book[config["sheet"]]
        expected = {"A1": "index", "F1": "DAC_A_nominal_V", "H1": "Vout_A_nominal_V"}
        header = {cell.coordinate: cell.value for row in sheet.iter_rows(max_row=1)
                  for cell in row if cell.value is not None}
        for cell, label in expected.items():
            if header.get(cell) != label:
                raise ValueError(f"入力形式が異なります: {cell}は{label}が必要です。")

        values, formulas = {}, {}
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    target = formulas if cell.data_type == "f" else values
                    target[cell.coordinate] = cell.value
        sheet_name = sheet.title
    finally:
        book.close()

    def number(cell):
        value = values.get(cell)
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            raise ValueError(f"{cell}: 有限な数値が必要です（値: {value!r}）。")
        return float(value)

    # 手順2–5: Vout=-H列、Ipol=-F列/10^6。電流はA、電圧はV。
    # 元Excelと行番号を合わせるため、解析行は生データ行+1（O3=-H2）。
    raw_rows = sorted(int(cell[1:]) for cell in values
                      if cell.startswith("A") and cell[1:].isdigit() and int(cell[1:]) >= 2)
    all_rows = [row + 1 for row in raw_rows]
    voltage = np.array([-number(f"H{row}") for row in raw_rows])
    current = np.array([-number(f"F{row}") / RESISTANCE_OHM for row in raw_rows])
    total = current * RESISTANCE_OHM + voltage

    # 手順6–7: -10<Vin+Vout<10の厳密な範囲。両端の±10 Vは含めない。
    applicable = (-10 < total) & (total < 10)
    pair = config.get("analysis_rows")
    if not isinstance(pair, (list, tuple)) or len(pair) != 2:
        raise ValueError("analysis_rowsは開始行・終了行の整数で指定してください。")
    first, last = pair
    if any(type(value) is not int for value in pair) or first >= last:
        raise ValueError("analysis_rowsは開始行・終了行の整数で指定してください。")
    rows = np.asarray(all_rows)
    selected = (rows >= first) & (rows <= last)
    if rows[selected].tolist() != list(range(first, last + 1)):
        raise ValueError("解析範囲に生データの存在しない行があります。")
    if not np.all(applicable[selected]):
        raise ValueError("解析範囲に適用判定（-10 < Vin+Vout < 10）の外側の点があります。")
    return Measurement(sheet_name, values, formulas, all_rows, voltage, current,
                       total, applicable, selected)


def read_trace(source, config, *, details=False):
    """Vfや近似範囲が未選択のときも、GUIで補正後の点を表示できる。"""
    measurement = load_measurement(source, config)
    rows, x = measurement.rows, measurement.voltage
    trace = correct_trace(x, measurement.current, rows)
    return (rows, x, trace) if details else (rows, x, trace.log_current)
