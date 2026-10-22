"""球状LP解析の統合版。Python 3.10以上で、この1ファイルだけを実行できる。

python lp_calculation.py --config settings.json
python lp_calculation.py --config settings.json --output-dir results

必要ライブラリ: numpy（CSV入力）、openpyxl（Excel入力時のみ）。
読む順序: fit_line → correct_trace → calculate_parameters → calculate → main。
"""
from __future__ import annotations

import argparse
import csv
from dataclasses import asdict, dataclass
import io
import json
import math
from pathlib import Path
import sys
from typing import NamedTuple
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

import numpy as np

ROOT = Path(__file__).resolve().parent


# 物理定数と球半径の初期値。イオン近似の傾き・切片の固定値ではない。
DEFAULT_PARAMETERS = {
    "probe_radius_m": 0.02,
    "electron_mass_kg": 9.10939e-31,
    "boltzmann_J_per_K": 1.38e-23,
    "electron_charge_C": 1.6e-19,
}
RESISTANCE_OHM = 1e6
ROUNDOFF_FACTOR = 32 * np.finfo(float).eps


def normalize_config(config):
    """現在の解析条件にそろえる。近似係数や領域を自動で決める処理はしない。"""
    settings = DEFAULT_PARAMETERS | dict(config)
    for key in ("ion_slope_A_per_V", "ion_intercept_A", "ion_fit_rows", "floating_potential_V"):
        settings.pop(key, None)
    settings.update(analysis_protocol="lp24", resistance_ohm=RESISTANCE_OHM,
                    ies_method="vp_neighbors", vp_fit_rows=None, floating_potential_rows=None)
    return settings


@dataclass(frozen=True)
class Fit:
    """y = slope*x + intercept と、使用した行・点数・決定係数R²。"""
    slope: float
    intercept: float
    r_squared: float | None
    rows: tuple[int, int]
    points: int

    def at(self, x):
        return self.slope * x + self.intercept


class Trace(NamedTuple):
    """イオン補正後の配列。NaNはln・差分を定義できない点を表す。"""
    ion_current: np.ndarray
    electron_current: np.ndarray
    log_current: np.ndarray
    derivative: np.ndarray
    ion_fit: Fit
    peak_index: int


def fit_line(x, y, rows) -> Fit:
    """切片付き最小二乗法。ExcelのSLOPE・INTERCEPTと同じ計算。"""
    x, y = np.asarray(x, dtype=float), np.asarray(y, dtype=float)
    if len(x) < 2 or not np.all(np.isfinite(x)) or not np.all(np.isfinite(y)):
        raise ValueError(f"近似範囲{rows}: 有限な点が2点以上必要です。")

    # a = Σ[(x-x平均)(y-y平均)] / Σ[(x-x平均)^2]、b = y平均-a*x平均。
    # @ は内積。途中の傾き・切片は丸めず、以降の計算にそのまま渡す。
    dx = x - x.mean()
    denominator = float(dx @ dx)
    if denominator == 0:
        raise ValueError(f"近似範囲{rows}: 電圧がすべて同一です。")
    slope = float(dx @ (y - y.mean()) / denominator)
    intercept = float(y.mean() - slope * x.mean())
    residual = float(np.sum((y - (slope * x + intercept)) ** 2))
    total = float(np.sum((y - y.mean()) ** 2))
    return Fit(slope, intercept, 1 - residual / total if total else None,
               tuple(rows), len(x))


def intersect_fits(retarding: Fit, saturation: Fit) -> float:
    """減速直線 ar*V+br と飽和直線 as*V+bs の交点の電圧Vp。"""
    denominator = retarding.slope - saturation.slope
    if math.isclose(denominator, 0, abs_tol=1e-14):
        raise ValueError("減速領域と飽和領域の直線が平行で、Vpを決定できません。")
    vp = (saturation.intercept - retarding.intercept) / denominator
    if not math.isfinite(vp):
        raise ValueError("交点のVpが有限値になりません。")
    return vp


def fit_vp_neighbors(voltage, logarithm, rows, vp) -> Fit:
    """正電圧の有効点からVp直前・直後を自動選択し、ln値を局所近似。"""
    x, y, rows = map(np.asarray, (voltage, logarithm, rows))
    valid = (x > 0) & np.isfinite(x) & np.isfinite(y)
    lo, hi = vp_bracketing_rows(x, np.where(valid, y, np.nan), rows, vp)
    mask = (rows >= lo) & (rows <= hi) & valid
    # Vpと同じ電圧の点があれば、それも含む。両側に点がなければ外挿しない。
    return fit_line(x[mask], y[mask], (lo, hi))


def correct_trace(x, y, rows) -> Trace:
    """負電圧のイオン近似 → 差し引き → ln変換 → Vp0用の差分。"""
    # 手順8–10: 適用範囲の負電圧点を全て使い、Iion=a*Vout+bを計算。
    # maskは選ぶ点を示すTrue/Falseの配列、x[mask]は選ばれた電圧だけの配列。
    mask = ion_region(x)
    selected = rows[mask]
    ion_fit = fit_line(x[mask], y[mask], (int(selected[0]), int(selected[-1])))
    a, b = ion_fit.slope, ion_fit.intercept
    ion = a * x + b
    electron = y - ion

    # 2点を通る直線の差し引きでは、本来0の残差が10^-23 A程度になることがある。
    # ln≈-50という丸め誤差を防ぐため、差し引く値の大きさと計算精度で0を判定する。
    # 固定の電流閾値やlnの下限は設けず、実際に小さい電流は保持する。
    tolerance = ROUNDOFF_FACTOR * (np.abs(y) + np.abs(a * x) + abs(b))
    electron = np.where(np.abs(electron) <= tolerance, 0., electron)
    if not np.all(np.isfinite(electron)):
        raise ValueError("補正電流に非有限値があります。")

    # 手順11–13: A単位の補正電流の絶対値を自然対数に変換する。
    # Ie=0のlnは未定義なのでNaNにし、グラフ・近似・Vf候補から除く。
    with np.errstate(divide="ignore"):
        logarithm = np.log(np.abs(electron))
    logarithm[~np.isfinite(logarithm)] = np.nan
    if not np.any(np.isfinite(logarithm)):
        raise ValueError("ln表示可能な補正電流がありません。")

    # 手順15: Vp0にはln変換前のIeを使う。np.diffは隣接点の差。
    # 最大区間の先頭点を採用し、次の点がない末尾はNaNにする。
    delta_x = np.diff(x)
    if np.any(delta_x == 0):
        raise ValueError("隣接する解析点の電圧差が0です。微分できません。")
    derivative = np.append(np.abs(np.diff(electron) / delta_x), np.nan)
    peak = int(np.nanargmax(derivative))
    return Trace(ion, electron, logarithm, derivative, ion_fit, peak)


def calculate_parameters(x, rows, trace: Trace, config):
    """手動選択した減速・飽和領域からVp・Ies・Te・電子密度を計算。"""
    parameters = {key: float(config.get(key, value))
                  for key, value in DEFAULT_PARAMETERS.items()}
    for key, value in parameters.items():
        if not math.isfinite(value) or value <= 0:
            raise ValueError(f"設定 {key} は正の有限値にしてください。")

    # 手順14: Vfは正電圧のln最小点、またはln値が小さい2点の電圧平均。
    # 明瞭な最小点かどうかは利用者がグラフを見て方法を選ぶ。
    mode = config.get("floating_potential_mode")
    if mode not in ("minimum", "lowest_two"):
        raise ValueError("Vfを目視で確認し、最小実測点か小さい2点の平均を選んでください。")
    floating = minimum_floating_potential(
        x, trace.log_current, rows, average_two=mode == "lowest_two", positive_only=True)

    # 手順16–19: ドラッグして保存した行範囲を使ってln(Ie)を近似する。
    # 減速領域はVf以上Vp0未満。飽和領域にはVp0による範囲制限を設けない。
    fits = {"ion": trace.ion_fit}
    for key, label in (("retarding", "電子減速領域"), ("saturation", "電子飽和領域")):
        pair = config.get(f"{key}_fit_rows")
        if not isinstance(pair, (list, tuple)) or len(pair) != 2:
            raise ValueError(f"{label}の開始・終了行を指定してください。")
        lo, hi = pair
        if not rows[0] <= lo < hi <= rows[-1]:
            raise ValueError(f"{label}は解析範囲内で2点以上選択してください。")
        mask = (rows >= lo) & (rows <= hi)
        if np.any(x[mask] <= 0):
            raise ValueError(f"{label}はVout > 0 Vの測定点から選んでください。")
        if key == "retarding":
            check_region_bounds(x, rows, pair, label=label,
                                lower=floating.voltage_V, upper=x[trace.peak_index])
        mask &= np.isfinite(trace.log_current)
        fits[key] = fit_line(x[mask], trace.log_current[mask], pair)

    # 手順20–22: 2直線の交点をVpとし、その前後の点の近似は自動で求める。
    # 局所直線へのVp代入値はln(Ies)。expで電流[A]に戻してから密度に使う。
    retarding = fits["retarding"]
    vp = intersect_fits(retarding, fits["saturation"])
    fits["vp"] = fit_vp_neighbors(x, trace.log_current, rows, vp)
    ln_ies = fits["vp"].at(vp)
    ies = math.exp(ln_ies)

    # 手順18: ln(Ie)の減速領域の傾きからTe[K]=e/(k*a_ret)を求める。
    # eは電気素量[C]。Te[eV]は1/a_ret。傾きが0以下なら計算を止める。
    if retarding.slope <= 0:
        raise ValueError("電子減速領域の傾きが0以下のため、正の電子温度を計算できません。")
    charge = parameters["electron_charge_C"]
    k = parameters["boltzmann_J_per_K"]
    mass = parameters["electron_mass_kg"]
    temperature = charge / k / retarding.slope

    # 手順23–24: 球面積S=4*pi*r^2、平均速度v=sqrt(8*k*Te/(pi*me))。
    # ne=4*Ies/(e*S*v)はm^-3。10^6で割るとcm^-3になる。
    area = 4 * math.pi * parameters["probe_radius_m"] ** 2
    speed = math.sqrt(8 * k * temperature / (math.pi * mass))
    density = 4 * ies / (charge * area * speed)
    metrics = {
        "probe_area_m2": area,
        "electron_temperature_K": temperature,
        "electron_temperature_eV": 1 / retarding.slope,
        "electron_mean_speed_m_per_s": speed,
        "plasma_potential_V": vp,
        "log_electron_saturation_current": ln_ies,
        "electron_saturation_current_A": ies,
        "electron_density_m3": density,
        "electron_density_cm3": density / 1e6,
        "density_log10_cm3": math.log10(density / 1e6),
        "derivative_peak_excel_row": int(rows[trace.peak_index]),
        "derivative_peak_voltage_V": float(x[trace.peak_index]),
        "floating_potential_V": floating.voltage_V,
        "minimum_log_voltage_V": floating.minimum_voltage_V,
    }
    if not all(math.isfinite(value) for value in metrics.values()):
        raise ValueError("結果に非有限値が生じました。測定値と近似範囲を確認してください。")
    return fits, metrics, floating


@dataclass(frozen=True)
class FloatingPotential:
    voltage_V: float
    rows: tuple[int, ...]
    endpoint_voltages_V: tuple[float, ...]
    minimum_row: int
    minimum_voltage_V: float
    minimum_log_current: float
    method: str

    @property
    def formula(self):
        if len(self.rows) == 1:
            return f'=AG{self.rows[0]}'
        return f"=(AG{self.rows[1]}+AG{self.rows[0]})/2"

    def to_dict(self):
        return {**asdict(self), 'excel_formula': self.formula}



def minimum_floating_potential(voltage, log_current, excel_rows, *, average_two=False,positive_only=False):
    """目視で方法を選んだ後、実測最小点またはln値が小さい2点からVfを計算。"""
    points = [(float(y), int(r), float(x)) for x,y,r in zip(voltage,log_current,excel_rows)
              if math.isfinite(float(x)) and math.isfinite(float(y)) and (not positive_only or float(x)>0)]
    points.sort(key=lambda point:(point[0],point[1]))
    count = 2 if average_two else 1
    if len(points) < count:
        raise ValueError('Vfを決定する有限なln表示の測定点が不足しています。'+(' Vout > 0 Vの点が必要です。' if positive_only else ''))
    chosen = points[:count]
    return FloatingPotential(sum(p[2] for p in chosen)/count,
                             tuple(p[1] for p in chosen), tuple(p[2] for p in chosen),
                             points[0][1], points[0][2], points[0][0],
                             'two_lowest_log_voltage_midpoint' if average_two else 'measured_log_minimum')



def ion_region(voltage):
    """表示中の適用範囲で、最小電圧から0 V未満までをすべて選ぶ。"""
    x = np.asarray(voltage, dtype=float)
    mask = np.isfinite(x) & (x < 0)
    if np.count_nonzero(mask) < 2:
        raise ValueError('イオン近似には、適用範囲内でVoutが0 V未満の測定点が2点以上必要です。')
    return mask


def check_region_bounds(voltage, rows, pair, *, label, lower, upper=None):
    x, rows = np.asarray(voltage), np.asarray(rows)
    lo, hi = pair
    mask = (rows >= lo) & (rows <= hi)
    valid = x[mask] >= lower
    if upper is not None:
        valid &= x[mask] < upper
    if not np.all(valid):
        bounds = f'{lower:.6g} V以上' + (f'、{upper:.6g} V未満' if upper is not None else '')
        raise ValueError(f'{label}は{bounds}の点から選択してください。')


def vp_bracketing_rows(voltage, logarithm, rows, vp):
    """Vpの両側にある最も近い有限な測定点で局所直線を作る。"""
    x, y, rows = np.asarray(voltage), np.asarray(logarithm), np.asarray(rows)
    valid = np.isfinite(y) & np.isfinite(x)
    below = np.flatnonzero(valid & (x < vp))
    above = np.flatnonzero(valid & (x > vp))
    if not len(below) or not len(above):
        raise ValueError('Vpの両側に測定点がありません。Iesを外挿せず、近似領域を見直してください。')
    left = int(below[np.argmax(x[below])])
    right = int(above[np.argmin(x[above])])
    # Vpと一致する測定点がある場合は、その点も含める。
    exact = np.flatnonzero(valid & (x == vp)).tolist()
    indices = sorted({left, right, *exact})
    skipped=set(range(min(indices),max(indices)+1))-set(indices)
    if any(valid[i] for i in skipped):
        raise ValueError('Vp近傍の行が連続していません。単一の掃引範囲を選択してください。')
    return [int(rows[min(indices)]), int(rows[max(indices)])]


@dataclass
class Calculation:
    """画面やExcel書式に依存しない計算結果。単位はV・A・K・m^-3。"""
    rows: np.ndarray
    voltage: np.ndarray
    current: np.ndarray
    trace: Trace
    fits: dict[str, Fit]
    metrics: dict[str, float]
    floating_potential: FloatingPotential


def load_config(path):
    """パスは設定JSONのあるフォルダを基準にするため、PCのユーザー名に依存しない。"""
    path = Path(path).resolve()
    if path.suffix.lower() != '.json':
        raise ValueError('設定JSONを指定してください。')
    config = json.loads(path.read_text(encoding='utf-8-sig'))
    if not isinstance(config, dict) or not all(key in config for key in ('input', 'analysis_rows')):
        raise ValueError('設定JSONにはinput・analysis_rowsが必要です。')
    return config


def choose_config(explicit=None, *, root=None):
    """指定がなければ、Vf方法と2領域を保存済みの最新設定を使う。"""
    if explicit is not None:
        return Path(explicit).resolve()
    root = Path(root or ROOT)
    candidates = [*root.joinpath('profiles').glob('*.json'),
                  *(root / name for name in ('settings.json', 'analysis_config.json', 'earth_analysis_config.json'))]
    ready = []
    for path in candidates:
        try:
            config = load_config(path)
            if (config.get('floating_potential_mode') in ('minimum', 'lowest_two')
                    and config.get('retarding_fit_rows') and config.get('saturation_fit_rows')
                    and (path.parent / config['input']).is_file()):
                ready.append(path)
        except (OSError, ValueError, KeyError, TypeError):
            continue
    if not ready:
        raise ValueError('計算用の保存済み設定がありません。--configでVf方法・減速領域・飽和領域を記入した設定JSONを指定してください。')
    return max(ready, key=lambda path: path.stat().st_mtime_ns).resolve()


def read_points(source, config):
    """LP512のCSV・Excelを直接読む。入力ファイルへの書き込みは行わない。"""
    source = Path(source)
    points = []

    def append(analysis_row, dac, vout):
        for value, label in ((dac, 'DAC_A_nominal_V'), (vout, 'Vout_A_nominal_V')):
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                raise ValueError(f'{analysis_row}行の{label}には有限な数値が必要です。')
        # 元Excelの解析列と同じ行番号: 生データのExcel行+1（O3=-H2）。
        points.append((analysis_row, -float(vout), -float(dac) / RESISTANCE_OHM))

    if source.suffix.lower() == '.csv':
        payload = source.read_bytes()
        for encoding in ('utf-8-sig', 'cp932', 'utf-16'):
            try:
                contents = payload.decode(encoding)
                break
            except UnicodeError:
                continue
        else:
            raise ValueError('CSVはUTF-8・Shift-JIS・UTF-16で保存してください。')
        try:
            dialect = csv.Sniffer().sniff(contents[:8192], delimiters=',\t;')
        except csv.Error:
            dialect = csv.excel
        records = csv.reader(io.StringIO(contents), dialect)
        header = [value.strip() for value in next(records, [])]
        required = ('DAC_A_nominal_V', 'Vout_A_nominal_V')
        if any(name not in header for name in required):
            raise ValueError('CSVにDAC_A_nominal_VとVout_A_nominal_Vの列が必要です。')
        indices = [header.index(name) for name in required]
        for record in records:
            if not any(value.strip() for value in record):
                continue
            row = len(points) + 3
            try:
                dac, vout = (float(record[index]) for index in indices)
            except (ValueError, IndexError) as error:
                raise ValueError(f'{row}行のDAC_A・Vout_Aには数値が必要です。') from error
            append(row, dac, vout)
    elif source.suffix.lower() == '.xlsx':
        import openpyxl  # CSVだけを使うPCでは不要。
        book = openpyxl.load_workbook(source, read_only=True, data_only=False)
        try:
            name = config.get('sheet') or book.sheetnames[0]
            if name not in book.sheetnames:
                raise ValueError(f'対象シートがありません: {name}')
            sheet = book[name]
            expected = {'A1': 'index', 'F1': 'DAC_A_nominal_V', 'H1': 'Vout_A_nominal_V'}
            for cell, label in expected.items():
                if sheet[cell].value != label:
                    raise ValueError(f'入力形式が異なります: {cell}は{label}が必要です。')
            for raw_row, values in enumerate(sheet.iter_rows(min_row=2, max_col=8, values_only=True), start=2):
                if values[0] is not None:
                    append(raw_row + 1, values[5], values[7])
        finally:
            book.close()
    else:
        raise ValueError('入力はLP512形式の.csvまたは.xlsxにしてください。')

    pair = config.get('analysis_rows')
    if (not isinstance(pair, (list, tuple)) or len(pair) != 2
            or any(type(value) is not int for value in pair) or pair[0] >= pair[1]):
        raise ValueError('analysis_rowsは開始行・終了行の整数で指定してください。')
    selected = [point for point in points if pair[0] <= point[0] <= pair[1]]
    if [point[0] for point in selected] != list(range(pair[0], pair[1] + 1)):
        raise ValueError('解析範囲に生データの存在しない行があります。')
    rows = np.array([point[0] for point in selected], dtype=int)
    x = np.array([point[1] for point in selected])
    y = np.array([point[2] for point in selected])
    total = y * RESISTANCE_OHM + x
    if not np.all((-10 < total) & (total < 10)):
        raise ValueError('解析範囲に適用判定（-10 < Vin+Vout < 10）の外側の点があります。')
    return rows, x, y


def calculate(source, config) -> Calculation:
    """生データから補正・ln・Vp・Te・密度までをまとめて計算する公開関数。"""
    config = normalize_config(config)
    rows, x, y = read_points(source, config)
    trace = correct_trace(x, y, rows)
    fits, metrics, floating = calculate_parameters(x, rows, trace, config)
    return Calculation(rows, x, y, trace, fits, metrics, floating)


def save_results(out, result, source, config, config_path):
    """持ち運び版では結果JSON・計算過程CSVだけを保存する。GUI・画像処理は不要。"""
    out = Path(out).resolve()
    summary, table = out / 'summary.json', out / 'analysis.csv'
    if {summary, table} & {Path(source).resolve(), Path(config_path).resolve()}:
        raise ValueError('出力先が入力・設定ファイルと重なっています。別フォルダを指定してください。')
    payload = {'input': str(source), 'config': config, 'metrics': result.metrics,
               'fits': {key: asdict(fit) for key, fit in result.fits.items()},
               'floating_potential': result.floating_potential.to_dict()}
    out.mkdir(parents=True, exist_ok=True)
    summary.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8')
    with table.open('w', newline='', encoding='utf-8-sig') as stream:
        writer = csv.writer(stream)
        writer.writerow(['excel_row', 'voltage_V', 'Ipol_A', 'ion_baseline_A', 'Ie_A', 'ln_abs_Ie', 'abs_dIe_dV'])
        arrays = (result.voltage, result.current, result.trace.ion_current,
                  result.trace.electron_current, result.trace.log_current, result.trace.derivative)
        for index, row in enumerate(result.rows):
            writer.writerow([int(row), *(float(a[index]) if np.isfinite(a[index]) else '' for a in arrays)])


def main(argv=None, *, root=None, prepare_source=None, exporter=None):
    """通常は表示のみ。既存プロジェクトの起動口は保存処理を渡して互換性を保つ。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config', type=Path, help='設定JSON。省略時は保存済みの最新設定')
    parser.add_argument('--input', type=Path, help='入力Excel/CSVを上書き指定')
    parser.add_argument('--output-dir', type=Path, help='指定時は結果JSON・CSVを保存')
    parser.add_argument('--no-export', action='store_true', help='結果を表示するだけで保存しない')
    args = parser.parse_args(argv)
    try:
        path = choose_config(args.config, root=root)
        config = normalize_config(load_config(path))
        source = (args.input or path.parent / config['input']).resolve()
        csv_input = source.suffix.lower() == '.csv'
        if prepare_source is not None:
            source = prepare_source(source)
            if csv_input and source.suffix.lower() == '.xlsx':
                config = config | {'sheet': 'LP_data'}
        result = calculate(source, config)
        print(f'使用設定: {path}\n入力: {source}')
        for key, label, unit in [
            ('floating_potential_V', 'Vf', 'V'),
            ('derivative_peak_voltage_V', 'Vp0', 'V'),
            ('plasma_potential_V', 'Vp', 'V'),
            ('electron_temperature_K', '電子温度', 'K'),
            ('electron_temperature_eV', '電子温度', 'eV'),
            ('electron_saturation_current_A', '電子飽和電流', 'A'),
            ('electron_density_cm3', '電子密度', 'cm^-3'),
        ]:
            print(f'{label}: {result.metrics[key]:.9g} {unit}')
        if not args.no_export and (args.output_dir is not None or exporter is not None):
            out = (args.output_dir or path.parent / config.get('output_dir', 'results')).resolve()
            if exporter is not None:
                exporter(source, config, out)
            else:
                save_results(out, result, source, config, path)
            print(f'保存先: {out}')
        return 0
    except (OSError, ValueError, KeyError, TypeError, OverflowError, ImportError, BadZipFile, ParseError) as error:
        print(f'解析エラー: {error}', file=sys.stderr)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
